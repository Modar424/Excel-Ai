"""End-to-end browser test: real FastAPI (AI_PROVIDER=mock) + production frontend build + Chromium (Playwright).
Run from repo root after `cd frontend && VITE_API_BASE_URL=http://localhost:8000 npm run build`:
    python e2e/run_e2e.py
Uses the deterministic mock AI (no Groq key/quota). Real-Groq behaviour is NOT covered here."""
import os
import subprocess
import sys
import tempfile
import time
import urllib.request
from io import BytesIO

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
API, WEB = "http://localhost:8000", "http://localhost:4173"
results: list[tuple[str, bool, str]] = []


def check(name, fn):
    try:
        fn()
        results.append((name, True, ""))
    except Exception as e:  # noqa: BLE001
        results.append((name, False, (str(e).splitlines() or [type(e).__name__])[0][:200]))


def wait_up(url, tries=40):
    for _ in range(tries):
        try:
            urllib.request.urlopen(url, timeout=1)
            return
        except Exception:  # noqa: BLE001
            time.sleep(0.5)
    raise RuntimeError(f"{url} did not start")


def main():
    env = {**os.environ, "AI_PROVIDER": "mock", "CORS_ORIGINS": WEB, "RATE_LIMIT_PER_MIN": "1000", "PYTHONPATH": os.path.join(ROOT, "backend")}
    be = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--port", "8000"], cwd=os.path.join(ROOT, "backend"), env=env,
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    fe = subprocess.Popen(["npx", "vite", "preview", "--port", "4173", "--strictPort"], cwd=os.path.join(ROOT, "frontend"), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shots = os.path.join(tempfile.gettempdir(), "e2e_shots"); os.makedirs(shots, exist_ok=True)
    try:
        wait_up(API + "/health"); wait_up(WEB)
        with sync_playwright() as p:
            b = p.chromium.launch()
            ctx = b.new_context(viewport={"width": 1280, "height": 900}, accept_downloads=True)
            pg = ctx.new_page()
            pg.set_default_timeout(20000)
            console = []
            pg.on("console", lambda m: console.append(m.text) if m.type == "error" else None)
            pg.goto(WEB); pg.wait_for_selector("#prompt")
            state = {}

            def ar_flow():
                assert pg.evaluate("document.documentElement.dir") == "rtl"
                pg.get_by_role("button", name="الموظفين", exact=True).click()
                assert "موظف" in pg.input_value("#prompt")
                pg.select_option("#rows", "3000")
                pg.click("button[type=submit]")
                expect(pg.get_by_text("تم إنشاء ملف Excel")).to_be_visible()
                expect(pg.locator("#result table th", has_text="الراتب")).to_be_visible()
                assert pg.locator("#result tbody tr").count() == 8
                expect(pg.get_by_text("الملف يحتوي على كل الصفوف")).to_be_visible()
                expect(pg.get_by_text("إجمالي الرواتب")).to_be_visible()
                pg.locator("#result").screenshot(path=f"{shots}/ar_result.png")
                with pg.expect_download() as dl:
                    pg.get_by_role("button", name="تنزيل الملف").click()
                d = dl.value
                state["ar_name"] = d.suggested_filename
                path = os.path.join(shots, "ar.xlsx"); d.save_as(path)
                wb = load_workbook(path)
                state["ar_sheets"] = wb.sheetnames
                assert wb.sheetnames == ["الموظفين", "الملخص"] and wb["الموظفين"].max_row == 3001 and wb["الموظفين"]["A3001"].value == 3000
                assert wb["الملخص"]["B4"].value.startswith("=")
            check("AR: example → 3,000-row generate → real xlsx download", ar_flow)
            check("AR: filename from Content-Disposition (UTF-8)", lambda: (_ for _ in ()).throw(AssertionError(state.get("ar_name"))) if state.get("ar_name") != "جدول_الموظفين.xlsx" else None)

            def en_flow():
                pg.get_by_role("button", name="English").first.click()
                assert pg.evaluate("document.documentElement.dir") == "ltr"
                pg.fill("#prompt", "Create an inventory sheet with product and quantity")
                pg.select_option("#rows", "500")
                pg.click("button[type=submit]")
                expect(pg.get_by_text("Your Excel file is ready")).to_be_visible()
                expect(pg.locator("#result table th", has_text="Quantity")).to_be_visible()
                with pg.expect_download() as dl:
                    pg.get_by_role("button", name="Download file").click()
                path = os.path.join(shots, "en.xlsx"); dl.value.save_as(path)
                wb = load_workbook(path)
                assert dl.value.suggested_filename == "Inventory.xlsx" and wb["Inventory"].max_row == 501 and wb.sheetnames == ["Inventory", "Summary"]
                pg.locator("#result").screenshot(path=f"{shots}/en_result.png")
            check("EN: LTR switch → 500-row generate → real xlsx download", en_flow)

            def err(trigger, snippet):
                def f():
                    pg.fill("#prompt", f"inventory {trigger}")
                    pg.click("button[type=submit]")
                    expect(pg.locator("[role=alert]", has_text=snippet)).to_be_visible()
                    expect(pg.get_by_role("button", name="Try again")).to_be_visible()
                return f
            check("EN error: AI quota shows message + retry hint", err("__quota__", "free AI quota"))
            check("EN error: quota includes seconds (42)", lambda: expect(pg.locator("div[role=alert]")).to_contain_text("about 42 seconds"))
            check("EN error: invalid spec", err("__invalid__", "could not be understood"))
            check("EN error: AI unavailable", err("__down__", "currently unavailable"))

            def net_fail():
                pg.route("**/api/v1/spreadsheets/generate", lambda r: r.abort())
                pg.fill("#prompt", "inventory sheet"); pg.click("button[type=submit]")
                expect(pg.locator("[role=alert]", has_text="Could not reach the server")).to_be_visible()
                pg.unroute("**/api/v1/spreadsheets/generate")
            check("EN error: network failure", net_fail)

            def proxy_html():
                pg.route("**/api/v1/spreadsheets/generate", lambda r: r.fulfill(status=502, content_type="text/html", body="<html>Bad gateway</html>"))
                pg.fill("#prompt", "inventory sheet"); pg.click("button[type=submit]")
                expect(pg.locator("[role=alert]", has_text="Could not create the file")).to_be_visible()
                pg.unroute("**/api/v1/spreadsheets/generate")
            check("EN error: non-JSON 502 from a proxy", proxy_html)

            def retry_works():
                pg.fill("#prompt", "inventory __down__"); pg.click("button[type=submit]")
                expect(pg.get_by_role("button", name="Try again")).to_be_visible()
                pg.fill("#prompt", "inventory sheet")
                pg.get_by_role("button", name="Try again").click()
                expect(pg.get_by_text("Your Excel file is ready")).to_be_visible()
            check("EN: retry after error succeeds", retry_works)

            def cancel():
                pg.fill("#prompt", "inventory __slow__"); pg.click("button[type=submit]")
                pg.get_by_role("button", name="Cancel").click()
                time.sleep(0.5)
                assert pg.locator("div[role=alert]").count() == 0
                expect(pg.locator("button[type=submit]")).to_be_enabled()
            check("EN: cancel in-flight request (no error shown)", cancel)

            def validation():
                pg.fill("#prompt", "ab")
                expect(pg.locator("button[type=submit]")).to_be_disabled()
                pg.fill("#prompt", "x" * 1001)
                expect(pg.get_by_text("Maximum 1000 characters")).to_be_visible()
                expect(pg.locator("button[type=submit]")).to_be_disabled()
                pg.fill("#prompt", "")
            check("Validation: short/over-long prompt disabled with message", validation)

            def theme():
                pg.get_by_role("switch").click()
                before = pg.evaluate("document.documentElement.classList.contains('dark')")
                pg.reload(); pg.wait_for_selector("#prompt")
                assert pg.evaluate("document.documentElement.classList.contains('dark')") == before
                assert pg.evaluate("localStorage.getItem('sc-lang')") == "en"
            check("Theme + language persist across reload", theme)
            ctx.close()

            m = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True, is_mobile=True, has_touch=True)
            mp = m.new_page(); mp.set_default_timeout(20000)

            def mobile():
                mp.goto(WEB); mp.wait_for_selector("#prompt")
                mp.get_by_role("button", name="المخزون", exact=True).click()
                mp.click("button[type=submit]")
                expect(mp.get_by_text("تم إنشاء ملف Excel")).to_be_visible()
                assert not mp.evaluate("document.documentElement.scrollWidth > innerWidth")
                with mp.expect_download() as dl:
                    mp.get_by_role("button", name="تنزيل الملف").click()
                p2 = os.path.join(shots, "m.xlsx"); dl.value.save_as(p2)
                assert load_workbook(p2)["المخزون"].max_row == 101  # default 100 rows
                mp.locator("#result").screenshot(path=f"{shots}/mobile_result.png")
            check("Mobile 390px (AR): generate + download, no horizontal overflow", mobile)
            m.close(); b.close()
            check("No browser console errors", lambda: (_ for _ in ()).throw(AssertionError(console[:2])) if [c for c in console if "Failed to load resource" not in c] else None)
    finally:
        for pr in (be, fe):
            pr.terminate()
    ok = all(r[1] for r in results)
    for n, good, msg in results:
        print(("PASS " if good else "FAIL ") + n + (f"  -> {msg}" if msg else ""))
    print(f"\n{sum(r[1] for r in results)}/{len(results)} passed; screenshots in {shots}")
    sys.exit(0 if ok else 1)


main()
