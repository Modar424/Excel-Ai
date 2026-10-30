from io import BytesIO

from openpyxl import load_workbook

from app.core.errors import AppError
from app.core.security import RateLimiter, safe_filename
from tests.conftest import MockProvider, make_client, sample_spec

URL = "/api/v1/spreadsheets/generate"
OK = {"prompt": "Create an employee sheet", "language": "en"}


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200 and r.json()["status"] == "ok"
    assert r.headers["x-request-id"]


def test_generate_returns_valid_xlsx(client):
    r = client.post(URL, json=OK)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="Employees.xlsx"' in r.headers["content-disposition"]
    assert r.headers["x-rows-generated"] == "2"
    ws = load_workbook(BytesIO(r.content))["Data"]
    assert [c.value for c in ws[1]] == ["Name", "Salary"]
    assert ws["B3"].value == 200


def test_formula_like_text_stays_text(client):
    ws = load_workbook(BytesIO(client.post(URL, json=OK).content))["Data"]
    assert ws["A3"].value == "=1+1" and ws["A3"].data_type == "s"


def test_arabic_prompt_and_title():
    c = make_client(MockProvider(sample_spec("جدول الموظفين")))
    r = c.post(URL, json={"prompt": "أنشئ جدولًا للموظفين", "language": "ar"})
    assert r.status_code == 200 and "filename*=UTF-8''" in r.headers["content-disposition"]


def test_invalid_requests(client):
    for body in [{}, {"prompt": ""}, {"prompt": "ab"}, {"prompt": "   "}, {"prompt": "x" * 4001},
                 {"prompt": "hello", "language": "fr"}, {"prompt": "hello", "max_rows": 0},
                 {"prompt": "hello", "max_rows": 5001}, {"prompt": "hello", "max_rows": "abc"}]:
        r = client.post(URL, json=body)
        assert r.status_code == 400, body
        e = r.json()["error"]
        assert e["code"] == "invalid_request" and e["request_id"]


def test_validation_error_does_not_echo_input(client):
    r = client.post(URL, json={"prompt": "hello", "language": "SECRET-VALUE"})
    assert "SECRET-VALUE" not in r.text


def test_oversize_body_413():
    c = make_client(MockProvider(sample_spec()), max_request_bytes=200)
    r = c.post(URL, json={"prompt": "x" * 500})
    assert r.status_code == 413 and r.json()["error"]["code"] == "payload_too_large"


def test_missing_key_503():
    c = make_client(None, groq_api_key="")
    r = c.post(URL, json=OK)
    assert r.status_code == 503 and r.json()["error"]["code"] == "ai_unavailable"


def test_ai_timeout_504():
    c = make_client(MockProvider(sample_spec(), delay=1), ai_total_timeout_s=0.05)
    assert c.post(URL, json=OK).status_code == 504


def test_provider_error_passthrough():
    c = make_client(MockProvider(exc=AppError(429, "ai_quota", "quota")))
    r = c.post(URL, json=OK)
    assert r.status_code == 429 and r.json()["error"]["code"] == "ai_quota"


def test_unexpected_error_is_safe_500():
    c = make_client(MockProvider(exc=RuntimeError("secret internals")))
    r = c.post(URL, json=OK)
    assert r.status_code == 500 and "secret" not in r.text and r.json()["error"]["code"] == "internal_error"


def test_rate_limit_429():
    c = make_client(MockProvider(sample_spec()), rate_limit_per_min=2)
    assert [c.post(URL, json=OK).status_code for _ in range(3)] == [200, 200, 429]
    r = c.post(URL, json=OK)
    assert int(r.headers["retry-after"]) >= 1 and r.json()["error"]["retry_after"] >= 1


def test_cors_allowed_and_blocked(client):
    ok = client.options(URL, headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST",
                                      "Access-Control-Request-Headers": "content-type"})
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:5173"
    bad = client.options(URL, headers={"Origin": "http://evil.example", "Access-Control-Request-Method": "POST"})
    assert "access-control-allow-origin" not in bad.headers
    r = client.get("/health", headers={"Origin": "http://evil.example"})
    assert "access-control-allow-origin" not in r.headers


def test_404_and_405_consistent(client):
    assert client.get("/nope").json()["error"]["code"] == "not_found"
    assert client.get(URL).json()["error"]["code"] == "method_not_allowed"


def test_openapi_available(client):
    assert URL in client.get("/openapi.json").json()["paths"]


def test_rate_limiter_unit():
    t = [0.0]
    rl = RateLimiter(2, clock=lambda: t[0])
    assert rl.check("a") == 0 and rl.check("a") == 0 and rl.check("a") > 0 and rl.check("b") == 0
    t[0] = 61
    assert rl.check("a") == 0


def test_safe_filename():
    assert safe_filename("../../etc/passwd")[0] == "etcpasswd.xlsx"
    assert safe_filename("")[0] == "workbook.xlsx"
    assert safe_filename('a"b;c\r\nd')[0] == "abcd.xlsx"
