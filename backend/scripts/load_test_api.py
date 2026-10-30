"""HTTP-level load test against a real uvicorn process (AI_PROVIDER=mock, so AI latency is excluded).
Measures request latency and the server's peak RSS (/proc VmHWM, Linux only). Usage: PYTHONPATH=. python scripts/load_test_api.py"""
import os
import statistics
import subprocess
import sys
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import httpx
from openpyxl import load_workbook

PORT = 8011
URL = f"http://localhost:{PORT}/api/v1/spreadsheets/generate"


def vm_hwm_mb(pid):
    for line in open(f"/proc/{pid}/status"):
        if line.startswith("VmHWM"):
            return round(int(line.split()[1]) / 1024)


def post(prompt, rows):
    t = time.perf_counter()
    r = httpx.post(URL, json={"prompt": prompt, "language": "en", "max_rows": rows}, timeout=180)
    return r.status_code, time.perf_counter() - t, len(r.content), r


def pct(xs, p):
    xs = sorted(xs)
    return xs[min(len(xs) - 1, int(len(xs) * p))]


def main():
    env = {**os.environ, "AI_PROVIDER": "mock", "RATE_LIMIT_PER_MIN": "100000", "MAX_ROWS": "5000", "MAX_CONCURRENT_BUILDS": os.environ.get("MAX_CONCURRENT_BUILDS", "2")}
    srv = subprocess.Popen(  # noqa: S603 (fixed argv)
        [sys.executable, "-m", "uvicorn", "app.main:app", "--port", str(PORT)], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        for _ in range(60):
            try:
                urllib.request.urlopen(f"http://localhost:{PORT}/health", timeout=1)
                break
            except Exception:  # noqa: BLE001
                time.sleep(0.5)
        print(f"server idle peak RSS: {vm_hwm_mb(srv.pid)} MB  (MAX_CONCURRENT_BUILDS={env['MAX_CONCURRENT_BUILDS']})\n")
        # A: sequential 3,000-row workbooks
        res = [post("inventory sheet", 3000) for _ in range(10)]
        lat = [r[1] for r in res]
        wb = load_workbook(BytesIO(res[0][3].content))
        print(f"A) 10 sequential 3,000-row requests: statuses={set(r[0] for r in res)} p50={statistics.median(lat):.2f}s p95={pct(lat, .95):.2f}s "
              f"bytes={res[0][2]:,} rows_in_file={wb.worksheets[0].max_row - 1} | server peak RSS {vm_hwm_mb(srv.pid)} MB")
        # B: 8 concurrent worst-case workbooks (5,000 rows x 40 populated columns)
        t = time.perf_counter()
        with ThreadPoolExecutor(8) as ex:
            res = list(ex.map(lambda _: post("__wide__", 5000), range(8)))
        wall = time.perf_counter() - t
        lat = [r[1] for r in res]
        print(f"B) 8 concurrent 5,000x40 requests: statuses={sorted(set(r[0] for r in res))} wall={wall:.1f}s p50={statistics.median(lat):.1f}s max={max(lat):.1f}s "
              f"bytes={res[0][2]:,} | server peak RSS {vm_hwm_mb(srv.pid)} MB")
        # C: 40 concurrent small requests
        t = time.perf_counter()
        with ThreadPoolExecutor(40) as ex:
            res = list(ex.map(lambda _: post("inventory sheet", 100), range(40)))
        lat = [r[1] for r in res]
        print(f"C) 40 concurrent 100-row requests: statuses={sorted(set(r[0] for r in res))} wall={time.perf_counter() - t:.1f}s p50={statistics.median(lat):.2f}s p95={pct(lat, .95):.2f}s "
              f"| server peak RSS {vm_hwm_mb(srv.pid)} MB")
    finally:
        srv.terminate()


main()
