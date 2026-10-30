"""Bugs found by fuzzing/QA in Phase 7 — kept as permanent regression tests."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.config import Settings
from app.schemas.spreadsheet import CalcSpec, ColumnSpec, SheetSpec, WorkbookSpec
from app.services.excel_service import build_workbook
from app.services.spreadsheet_service import validate_spec

S = Settings(_env_file=None)


def build(columns, rows, **kw):
    sh = SheetSpec(title=kw.pop("title", "S"), columns=columns, rows=rows, **kw)
    return load_workbook(BytesIO(build_workbook(validate_spec(WorkbookSpec(title="T", sheets=[sh]), S, 100), S).data))


@pytest.mark.parametrize("bad", ["\x00", "a\x08b", "x\x1fy", "\x0b"])
@pytest.mark.parametrize("ctype", ["text", "integer", "date", "boolean"])
def test_xml_illegal_control_chars_never_crash_any_column_type(bad, ctype):
    wb = build([ColumnSpec(name="c", type=ctype)], [[bad]])
    assert "\x00" not in str(wb["S"]["A2"].value) and "\x08" not in str(wb["S"]["A2"].value)


def test_control_chars_in_headers_labels_notes_and_titles():
    sh = SheetSpec(title="a\x00b", columns=[ColumnSpec(name="h\x08x")], rows=[["v"]], notes=["n\x01"],
                   calculations=[CalcSpec(label="l\x02", kind="counta", column="h x")])
    wb = load_workbook(BytesIO(build_workbook(validate_spec(WorkbookSpec(title="t\x03", sheets=[sh]), S, 10), S).data))
    assert wb.sheetnames[0] == "a b" and wb["a b"]["A1"].value == "h x"


@pytest.mark.parametrize("v,expected_text", [(1.7976931348623155e308, True), (10**400, True), ("1e999", True), ("-1e999", True), (1e20, False)])
def test_out_of_range_numbers_become_text_not_invalid_cells(v, expected_text):
    wb = build([ColumnSpec(name="n", type="decimal")], [[v]])
    cell = wb["S"]["A2"]
    assert (cell.data_type == "s") is expected_text
