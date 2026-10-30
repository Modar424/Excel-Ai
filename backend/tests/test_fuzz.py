"""Property-based tests: adversarial AI specs and request bodies must never crash the server, never create
formulas outside the validated Summary column, and always produce a readable workbook or a clean AppError."""
import math
import os
from io import BytesIO

from hypothesis import HealthCheck, assume, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.core.config import Settings
from app.core.errors import AppError
from app.schemas.spreadsheet import WorkbookSpec
from app.services.excel_service import build_workbook
from app.services.formula_policy import validate_formula
from app.services.mock_provider import MockProvider
from app.services.spreadsheet_service import validate_spec
from tests.conftest import make_client

NASTY = ["=1+1", "+cmd|' /C calc'!A0", "@SUM(1)", "-1+2", "\t=x", '=HYPERLINK("http://x","y")', "'", '"', "[", "]", "*", "?", "~", "\x00",
         "\u202e", "a" * 400, "=" * 50, "', '!A1", "الاسم", "😀", "", " ", "0", "-0", "1e999", "NaN", ">=", "<>", '"; =1+1', "Sheet1'!A1", "\n=1"]
txt = st.one_of(st.sampled_from(NASTY), st.text(max_size=40))
cell = st.one_of(st.none(), st.booleans(), st.integers(-10**12, 10**12), st.floats(allow_nan=True, allow_infinity=True), txt)
TYPES = ["text", "integer", "decimal", "currency", "percent", "date", "boolean"]
KINDS = ["count", "counta", "sum", "average", "min", "max", "countif", "sumif", "averageif", "countifs", "ratio", "unique_count"]
GEN = ["sequence", "constant", "date_sequence", "cycle"]


@st.composite
def specs(draw):
    ncols = draw(st.integers(1, 6))
    names = [draw(txt) or "c" for _ in range(ncols)]
    pick = st.one_of(st.sampled_from(names), txt, st.none())
    cols = [{"name": n, "type": draw(st.sampled_from(TYPES))} for n in names]
    rows = [draw(st.lists(cell, min_size=ncols, max_size=ncols)) for _ in range(draw(st.integers(0, 6)))]
    rules = [{"column": draw(pick) or "x", "kind": draw(st.sampled_from(GEN)), "start": draw(st.one_of(st.none(), txt)),
              "step": draw(st.one_of(st.none(), st.floats(-1e6, 1e6), st.just(1e300))), "values": draw(st.lists(txt, max_size=3))} for _ in range(draw(st.integers(0, 3)))]
    calcs = [{"label": draw(txt) or "l", "kind": draw(st.sampled_from(KINDS)), "column": draw(pick), "criteria_column": draw(pick), "criteria": draw(st.one_of(st.none(), txt)),
              "group_by": draw(pick), "group_values": draw(st.lists(txt, max_size=3))} for _ in range(draw(st.integers(0, 5)))]
    sheets = [{"title": draw(txt) or "S", "columns": cols, "rows": rows, "row_count": draw(st.integers(0, 40)), "generation_rules": rules,
               "calculations": calcs, "notes": draw(st.lists(txt, max_size=2)), "synthetic": draw(st.booleans())} for _ in range(draw(st.integers(1, 2)))]
    return {"title": draw(txt) or "W", "sheets": sheets, "rtl": draw(st.booleans())}


@settings(max_examples=int(os.getenv("FUZZ_EXAMPLES", "120")), deadline=None, suppress_health_check=[HealthCheck.too_slow, HealthCheck.data_too_large])
@given(specs(), st.sampled_from([10**9, 1]))
def test_adversarial_specs_are_safe(raw, threshold):
    try:
        spec = WorkbookSpec.model_validate(raw)
    except ValueError:
        assume(False)
    s = Settings(_env_file=None, writeonly_cells_threshold=threshold)
    try:
        spec = validate_spec(spec, s, 100)
    except AppError as e:
        assert e.status == 422
        return
    res = build_workbook(spec, s)  # must not raise
    wb = load_workbook(BytesIO(res.data))
    names = {sh.title for sh in spec.sheets}
    summary = wb.sheetnames[len(spec.sheets):]
    for ws in wb.worksheets:
        is_summary = ws.title in summary
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    assert is_summary and c.column == 2, f"formula outside Summary!B at {ws.title}!{c.coordinate}"
                    validate_formula(c.value, names)
                if isinstance(c.value, float):
                    assert math.isfinite(c.value)
    assert all(len(ws.title) <= 31 for ws in wb.worksheets)


json_leaf = st.one_of(st.none(), st.booleans(), st.integers(-10**20, 10**20), st.floats(allow_nan=False, allow_infinity=False), st.text(max_size=60))
json_val = st.recursive(json_leaf, lambda c: st.one_of(st.lists(c, max_size=4), st.dictionaries(st.text(max_size=8), c, max_size=4)), max_leaves=12)
body = st.fixed_dictionaries({}, optional={"prompt": json_val, "language": json_val, "max_rows": json_val, "extra": json_val})


@settings(max_examples=int(os.getenv("FUZZ_EXAMPLES", "120")) + 30, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(body)
def test_random_request_bodies_never_500(b):
    c = make_client(MockProvider(), max_rows=50)
    r = c.post("/api/v1/spreadsheets/generate", json=b)
    assert r.status_code in (200, 400, 413, 422), (r.status_code, r.text[:200])
    if r.status_code != 200:
        e = r.json()["error"]
        assert set(e) >= {"code", "message", "request_id"} and "Traceback" not in r.text
